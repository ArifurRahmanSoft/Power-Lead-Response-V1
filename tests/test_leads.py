from io import BytesIO
import json
import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from app.core.config import settings
from app.core.permissions import PermissionKey, ROLE_PERMISSIONS, RoleKey
from app.imports.lead_excel import CANONICAL_COLUMNS
from app.models.lead import Lead, LeadCustomerAction, LeadEvent, LeadImportRow
from app.models.membership import TenantMembership
from app.models.role import Permission, Role, RolePermission
from app.models.tenant import Tenant
from app.models.user import User
from app.schemas.lead import LeadFields
from app.utils.password import hash_password

PASSWORD = "lead-test-password"


def seed_access(database_session: Session) -> dict[RoleKey, Role]:
    roles = {
        key: Role(
            key=key.value,
            name=key.value.title(),
            tenant_id=None,
            is_system=True,
        )
        for key in RoleKey
    }
    permissions = {
        key: Permission(key=key.value, description=key.value) for key in PermissionKey
    }
    database_session.add_all([*roles.values(), *permissions.values()])
    database_session.flush()
    database_session.add_all(
        RolePermission(
            role_id=roles[role_key].id,
            permission_id=permissions[permission].id,
        )
        for role_key, permission_keys in ROLE_PERMISSIONS.items()
        for permission in permission_keys
    )
    database_session.commit()
    return roles


def create_user(database_session: Session, login_id: str) -> User:
    user = User(
        name=login_id,
        login_id=login_id,
        email=f"{login_id}@example.com",
        password_hash=hash_password(PASSWORD),
    )
    database_session.add(user)
    database_session.commit()
    return user


def create_tenant(database_session: Session, slug: str) -> Tenant:
    tenant = Tenant(name=slug.title(), slug=slug)
    database_session.add(tenant)
    database_session.commit()
    return tenant


def add_membership(
    database_session: Session,
    user: User,
    tenant: Tenant,
    role: Role,
) -> TenantMembership:
    membership = TenantMembership(
        tenant_id=tenant.id,
        user_id=user.id,
        role_id=role.id,
    )
    database_session.add(membership)
    database_session.commit()
    return membership


def login_headers(client: TestClient, user: User) -> dict[str, str]:
    response = client.post(
        "/api/auth/login",
        json={"identifier": user.login_id, "password": PASSWORD},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['data']['access_token']}"}


def lead_payload(email: str, **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "first_name": "  Jane ",
        "last_name": " Doe  ",
        "email": email,
        "phone": "+88001234567",
        "country": "BD",
        "company_size": "11_50",
        "revenue_amount": "1000.00",
        "revenue_currency": "USD",
        "revenue_period": "annual",
    }
    payload.update(overrides)
    return payload


def workbook_bytes(rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    sheet.append(list(CANONICAL_COLUMNS))
    for row in rows:
        sheet.append([row.get(column) for column in CANONICAL_COLUMNS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def owner_setup(
    client: TestClient,
    database_session: Session,
) -> tuple[dict[RoleKey, Role], Tenant, User, dict[str, str]]:
    roles = seed_access(database_session)
    tenant = create_tenant(database_session, "lead-workspace")
    owner = create_user(database_session, "lead.owner")
    add_membership(database_session, owner, tenant, roles[RoleKey.OWNER])
    return roles, tenant, owner, login_headers(client, owner)


def test_lead_crud_duplicate_concurrency_and_archive(
    client: TestClient,
    database_session: Session,
) -> None:
    _, tenant, owner, headers = owner_setup(client, database_session)

    invalid = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload("invalid@example.com", first_name="   "),
    )
    assert invalid.status_code == 422

    created = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload("  Jane@Example.COM "),
    )
    assert created.status_code == 201
    lead = created.json()["data"]
    assert lead["first_name"] == "Jane"
    assert lead["email"] == "jane@example.com"
    assert lead["email_status"] == "unknown"
    assert lead["source"] == "manual"
    assert lead["tracking_id"].startswith("LD-")

    duplicate = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload("jane@example.com"),
    )
    assert duplicate.status_code == 409
    assert duplicate.json()["detail"]["existing_lead"]["id"] == lead["id"]

    filtered = client.get("/api/leads?search=Jane&country=BD", headers=headers)
    assert filtered.status_code == 200
    assert filtered.json()["data"]["total"] == 1

    invalid_transition = client.patch(
        f"/api/leads/{lead['id']}",
        headers=headers,
        json={"row_version": 1, "status": "converted"},
    )
    assert invalid_transition.status_code == 400

    updated = client.patch(
        f"/api/leads/{lead['id']}",
        headers=headers,
        json={"row_version": 1, "company": "New Company", "status": "qualified"},
    )
    assert updated.status_code == 200
    assert updated.json()["data"]["row_version"] == 2
    assert updated.json()["data"]["status"] == "qualified"

    stale = client.patch(
        f"/api/leads/{lead['id']}",
        headers=headers,
        json={"row_version": 1, "company": "Stale Company"},
    )
    assert stale.status_code == 409

    action = LeadCustomerAction(
        tenant_id=tenant.id,
        lead_id=uuid.UUID(lead["id"]),
        action_type="send_email",
        status="pending",
    )
    database_session.add(action)
    database_session.commit()
    archived = client.delete(
        f"/api/leads/{lead['id']}?row_version=2",
        headers=headers,
    )
    assert archived.status_code == 200
    database_session.expire_all()
    assert database_session.get(LeadCustomerAction, action.id).status == "cancelled"
    assert client.get("/api/leads", headers=headers).json()["data"]["total"] == 0
    archived_lead = database_session.get(Lead, uuid.UUID(lead["id"]))
    assert archived_lead.deleted_at is not None
    assert archived_lead.deleted_by == owner.id
    assert database_session.scalar(
        select(LeadEvent).where(LeadEvent.event_type == "lead.archived")
    ) is not None


def test_assignee_validation_tenant_isolation_and_staff_scope(
    client: TestClient,
    database_session: Session,
) -> None:
    roles, tenant, _, owner_headers = owner_setup(client, database_session)
    staff = create_user(database_session, "assigned.staff")
    other_staff = create_user(database_session, "other.staff")
    outsider = create_user(database_session, "outside.user")
    add_membership(database_session, staff, tenant, roles[RoleKey.STAFF])
    add_membership(database_session, other_staff, tenant, roles[RoleKey.STAFF])

    invalid_assignee = client.post(
        "/api/leads",
        headers=owner_headers,
        json=lead_payload("invalid-assignee@example.com", assigned_user_id=str(outsider.id)),
    )
    assert invalid_assignee.status_code == 400

    assigned = client.post(
        "/api/leads",
        headers=owner_headers,
        json=lead_payload("assigned@example.com", assigned_user_id=str(staff.id)),
    ).json()["data"]
    unassigned_to_staff = client.post(
        "/api/leads",
        headers=owner_headers,
        json=lead_payload("other@example.com", assigned_user_id=str(other_staff.id)),
    ).json()["data"]

    staff_headers = login_headers(client, staff)
    staff_list = client.get("/api/leads", headers=staff_headers)
    assert staff_list.status_code == 200
    assert [item["id"] for item in staff_list.json()["data"]["items"]] == [assigned["id"]]
    assert client.get(f"/api/leads/{unassigned_to_staff['id']}", headers=staff_headers).status_code == 403

    staff_update = client.patch(
        f"/api/leads/{assigned['id']}",
        headers=staff_headers,
        json={"row_version": 1, "notes": "Assigned update"},
    )
    assert staff_update.status_code == 200
    staff_status_change = client.patch(
        f"/api/leads/{assigned['id']}",
        headers=staff_headers,
        json={"row_version": 2, "status": "qualified"},
    )
    assert staff_status_change.status_code == 403

    other_tenant = create_tenant(database_session, "other-tenant")
    other_owner = create_user(database_session, "other.owner")
    add_membership(database_session, other_owner, other_tenant, roles[RoleKey.OWNER])
    other_headers = login_headers(client, other_owner)
    assert client.get(f"/api/leads/{assigned['id']}", headers=other_headers).status_code == 404
    same_email_other_tenant = client.post(
        "/api/leads",
        headers=other_headers,
        json=lead_payload("assigned@example.com"),
    )
    assert same_email_other_tenant.status_code == 201


def test_excel_preview_commit_formulas_invalid_rows_and_retry(
    client: TestClient,
    database_session: Session,
) -> None:
    _, _, _, headers = owner_setup(client, database_session)
    existing = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload("duplicate@example.com"),
    )
    assert existing.status_code == 201

    file_bytes = workbook_bytes(
        [
            lead_payload("new@example.com"),
            lead_payload("duplicate@example.com"),
            lead_payload("not-an-email"),
            lead_payload("formula@example.com", first_name="=1+1"),
        ]
    )
    preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={
            "file": (
                "leads.xlsx",
                file_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 201
    batch = preview.json()["data"]
    assert batch["summary"] == {
        "total": 4,
        "valid": 1,
        "invalid": 2,
        "duplicate": 1,
        "inserted": 0,
        "failed": 0,
    }
    assert database_session.scalar(select(Lead).where(Lead.email == "new@example.com")) is None

    errors = client.get(f"/api/leads/imports/{batch['id']}/errors", headers=headers)
    assert errors.status_code == 200
    assert errors.json()["data"]["total"] == 3

    committed = client.post(f"/api/leads/imports/{batch['id']}/commit", headers=headers)
    assert committed.status_code == 200
    assert committed.json()["data"]["summary"]["inserted"] == 1
    assert database_session.scalar(select(Lead).where(Lead.email == "new@example.com")) is not None

    repeated = client.post(f"/api/leads/imports/{batch['id']}/commit", headers=headers)
    assert repeated.status_code == 200
    assert database_session.scalars(select(Lead).where(Lead.email == "new@example.com")).all().__len__() == 1

    same_preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={"file": ("leads.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert same_preview.status_code == 201
    assert same_preview.json()["data"]["id"] == batch["id"]

    mapped_workbook = Workbook()
    mapped_sheet = mapped_workbook.active
    mapped_sheet.append(["Given Name", "Family Name", "Email Address"])
    mapped_sheet.append(["Mapped", "Person", "mapped@example.com"])
    mapped_output = BytesIO()
    mapped_workbook.save(mapped_output)
    mapped_preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={"file": ("mapped.xlsx", mapped_output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={
            "mapping": json.dumps(
                {
                    "Given Name": "first_name",
                    "Family Name": "last_name",
                    "Email Address": "email",
                }
            )
        },
    )
    assert mapped_preview.status_code == 201
    assert mapped_preview.json()["data"]["summary"]["valid"] == 1


def test_import_partial_failure_limits_template_and_permissions(
    client: TestClient,
    database_session: Session,
    monkeypatch,
) -> None:
    roles, tenant, _, headers = owner_setup(client, database_session)
    template = client.get("/api/leads/import-template", headers=headers)
    assert template.status_code == 200
    workbook = load_workbook(BytesIO(template.content), read_only=True)
    assert {"Instructions", "Leads", "Sample", "Catalogs"} <= set(workbook.sheetnames)
    assert "service_requested" in [cell.value for cell in workbook["Leads"][1]]
    sample_headers = [cell.value for cell in workbook["Sample"][1]]
    sample_values = [cell.value for cell in workbook["Sample"][2]]
    sample = dict(zip(sample_headers, sample_values, strict=True))
    assert sample["company_website"] == "google.com"
    assert sample["country"] == "Bangladesh"
    assert sample["service_requested"] == "Website Design"

    file_bytes = workbook_bytes(
        [lead_payload("first@example.com"), lead_payload("second@example.com")]
    )
    preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={"file": ("partial.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview.status_code == 201
    batch_id = uuid.UUID(preview.json()["data"]["id"])
    rows = database_session.scalars(
        select(LeadImportRow).where(LeadImportRow.batch_id == batch_id).order_by(LeadImportRow.row_number)
    ).all()
    rows[1].normalized_data = {"email": "broken@example.com"}
    flag_modified(rows[1], "normalized_data")
    database_session.commit()
    committed = client.post(f"/api/leads/imports/{batch_id}/commit", headers=headers)
    assert committed.status_code == 200
    assert committed.json()["data"]["status"] == "completed_with_errors"
    assert committed.json()["data"]["summary"]["inserted"] == 1
    assert committed.json()["data"]["summary"]["failed"] == 1

    monkeypatch.setattr(settings, "lead_import_max_rows", 1)
    too_many = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={"file": ("too-many.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert too_many.status_code == 400

    readonly = create_user(database_session, "lead.readonly")
    add_membership(database_session, readonly, tenant, roles[RoleKey.READONLY])
    readonly_headers = login_headers(client, readonly)
    assert client.get("/api/leads/import-template", headers=readonly_headers).status_code == 403


@pytest.mark.parametrize(
    ("raw_url", "normalized"),
    [
        ("google.com", "https://google.com"),
        ("www.google.com", "https://www.google.com"),
        ("https://google.com", "https://google.com"),
        ("http://google.com", "http://google.com"),
        (" google.com/path?q=lead ", "https://google.com/path?q=lead"),
    ],
)
def test_url_normalization_for_every_lead_url_field(
    raw_url: str,
    normalized: str,
) -> None:
    for field in (
        "company_website",
        "linkedin_url",
        "facebook_url",
        "instagram_url",
        "x_url",
        "github_url",
    ):
        lead = LeadFields.model_validate(
            {"first_name": "Jane", "last_name": "Doe", "email": "jane@example.com", field: raw_url}
        )
        assert getattr(lead, field) == normalized


@pytest.mark.parametrize(
    "invalid_url",
    [
        "javascript:alert(1)",
        "data:text/plain,test",
        "file:///tmp/test",
        "ftp://example.com/file",
        "https://user:password@example.com",
        "https://bad host.example.com",
        "not-a-host",
        "https://example.com:invalid",
    ],
)
def test_url_validation_rejects_unsafe_and_malformed_values(invalid_url: str) -> None:
    with pytest.raises(ValidationError):
        LeadFields.model_validate(
            {
                "first_name": "Jane",
                "last_name": "Doe",
                "email": "jane@example.com",
                "company_website": invalid_url,
            }
        )


def test_country_names_aliases_codes_and_unknown_values() -> None:
    examples = {
        "Bangladesh": "Bangladesh",
        " united states ": "United States",
        "Australia": "Australia",
        "United Kingdom": "United Kingdom",
        "UK": "United Kingdom",
        "usa": "United States",
        "BD": "Bangladesh",
        "USA": "United States",
    }
    for entered, canonical in examples.items():
        lead = LeadFields.model_validate(
            {"first_name": "Jane", "last_name": "Doe", "email": "jane@example.com", "country": entered}
        )
        assert lead.country == canonical

    for invalid_country in ("Atlantis", "Congo"):
        with pytest.raises(ValidationError):
            LeadFields.model_validate(
                {
                    "first_name": "Jane",
                    "last_name": "Doe",
                    "email": "jane@example.com",
                    "country": invalid_country,
                }
            )


def test_crud_and_excel_share_normalization_and_old_files_still_import(
    client: TestClient,
    database_session: Session,
) -> None:
    _, _, _, headers = owner_setup(client, database_session)
    created = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload(
            "normalized@example.com",
            country=" uk ",
            company_website=" google.com/path?q=1 ",
            linkedin_url="",
            service_requested="  Website Design  ",
        ),
    )
    assert created.status_code == 201
    lead = created.json()["data"]
    assert lead["country"] == "United Kingdom"
    assert lead["company_website"] == "https://google.com/path?q=1"
    assert lead["linkedin_url"] is None
    assert lead["service_requested"] == "Website Design"
    stored = database_session.get(Lead, uuid.UUID(lead["id"]))
    assert stored.country == "GB"

    updated = client.patch(
        f"/api/leads/{lead['id']}",
        headers=headers,
        json={"row_version": lead["row_version"], "service_requested": "   ", "country": "BD"},
    )
    assert updated.status_code == 200
    assert updated.json()["data"]["service_requested"] is None
    assert updated.json()["data"]["country"] == "Bangladesh"

    imported_file = workbook_bytes(
        [
            lead_payload(
                "import-normalized@example.com",
                country="United States",
                company_website="www.google.com",
                service_requested="  SEO  ",
            )
        ]
    )
    preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={"file": ("normalized.xlsx", imported_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview.status_code == 201
    batch_id = uuid.UUID(preview.json()["data"]["id"])
    import_row = database_session.scalar(select(LeadImportRow).where(LeadImportRow.batch_id == batch_id))
    assert import_row.normalized_data["country"] == "United States"
    assert import_row.normalized_data["company_website"] == "https://www.google.com"
    assert import_row.normalized_data["service_requested"] == "SEO"
    committed = client.post(f"/api/leads/imports/{batch_id}/commit", headers=headers)
    assert committed.status_code == 200
    imported = database_session.scalar(select(Lead).where(Lead.email == "import-normalized@example.com"))
    assert imported.country == "US"
    assert imported.company_website == "https://www.google.com"
    assert imported.service_requested == "SEO"

    old_workbook = Workbook()
    old_sheet = old_workbook.active
    old_sheet.title = "Leads"
    old_sheet.append(["first_name", "last_name", "email"])
    old_sheet.append(["Old", "Format", "old-format@example.com"])
    old_output = BytesIO()
    old_workbook.save(old_output)
    old_preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={"file": ("old.xlsx", old_output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert old_preview.status_code == 201
    assert old_preview.json()["data"]["summary"]["valid"] == 1
    old_commit = client.post(
        f"/api/leads/imports/{old_preview.json()['data']['id']}/commit",
        headers=headers,
    )
    assert old_commit.status_code == 200
    assert old_commit.json()["data"]["summary"]["inserted"] == 1


def test_service_length_and_import_country_errors(
    client: TestClient,
    database_session: Session,
) -> None:
    _, _, _, headers = owner_setup(client, database_session)
    too_long = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload("long-service@example.com", service_requested="x" * 201),
    )
    assert too_long.status_code == 422

    invalid_country = client.post(
        "/api/leads",
        headers=headers,
        json=lead_payload("unknown-country@example.com", country="Atlantis"),
    )
    assert invalid_country.status_code == 422
    assert invalid_country.json()["errors"][0]["field"] == "country"

    preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={
            "file": (
                "unknown-country.xlsx",
                workbook_bytes([lead_payload("import-country@example.com", country="Atlantis")]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 201
    assert preview.json()["data"]["summary"]["invalid"] == 1
    errors = client.get(
        f"/api/leads/imports/{preview.json()['data']['id']}/errors",
        headers=headers,
    )
    assert errors.status_code == 200
    assert errors.json()["data"]["items"][0]["errors"][0]["field"] == "country"


def test_service_filter_and_options_respect_scope_and_pagination(
    client: TestClient,
    database_session: Session,
) -> None:
    roles, tenant, _, owner_headers = owner_setup(client, database_session)
    staff = create_user(database_session, "service.staff")
    other_staff = create_user(database_session, "service.other")
    add_membership(database_session, staff, tenant, roles[RoleKey.STAFF])
    add_membership(database_session, other_staff, tenant, roles[RoleKey.STAFF])

    def create_service_lead(
        email: str,
        service: str | None,
        assignee: User,
        company: str,
    ) -> dict[str, object]:
        response = client.post(
            "/api/leads",
            headers=owner_headers,
            json=lead_payload(
                email,
                company=company,
                service_requested=service,
                assigned_user_id=str(assignee.id),
            ),
        )
        assert response.status_code == 201
        return response.json()["data"]

    website = create_service_lead(
        "service-web@example.com", "Website Design", staff, "Acme"
    )
    create_service_lead(
        "service-web-case@example.com", "website design", other_staff, "Beta"
    )
    create_service_lead("service-seo@example.com", "SEO", staff, "Acme")
    create_service_lead(
        "service-marketing@example.com", "Marketing", other_staff, "Acme"
    )
    create_service_lead("service-blank@example.com", "   ", staff, "Acme")
    archived = create_service_lead(
        "service-archived@example.com", "Archived Service", staff, "Acme"
    )
    archived_response = client.delete(
        f"/api/leads/{archived['id']}?row_version={archived['row_version']}",
        headers=owner_headers,
    )
    assert archived_response.status_code == 200

    combined = client.get(
        "/api/leads",
        headers=owner_headers,
        params={
            "service_requested": "  WEBSITE DESIGN  ",
            "company": "Acme",
            "page": 1,
            "size": 1,
        },
    )
    assert combined.status_code == 200
    assert combined.json()["data"]["total"] == 1
    assert [item["id"] for item in combined.json()["data"]["items"]] == [
        website["id"]
    ]

    paged_service = client.get(
        "/api/leads",
        headers=owner_headers,
        params={"service_requested": "website design", "page": 1, "size": 1},
    )
    assert paged_service.status_code == 200
    assert len(paged_service.json()["data"]["items"]) == 1
    assert paged_service.json()["data"]["total"] == 2

    no_substring = client.get(
        "/api/leads",
        headers=owner_headers,
        params={"service_requested": "website"},
    )
    assert no_substring.status_code == 200
    assert no_substring.json()["data"]["total"] == 0

    unfiltered = client.get(
        "/api/leads",
        headers=owner_headers,
        params={"service_requested": "   "},
    )
    assert unfiltered.status_code == 200
    assert unfiltered.json()["data"]["total"] == 5

    first_options = client.get(
        "/api/leads/service-requested-options",
        headers=owner_headers,
        params={"page": 1, "size": 2},
    )
    assert first_options.status_code == 200
    assert first_options.json() == {
        "success": True,
        "data": {
            "items": [
                {"label": "Marketing", "value": "marketing"},
                {"label": "SEO", "value": "seo"},
            ],
            "page": 1,
            "size": 2,
            "total": 3,
        },
    }
    second_options = client.get(
        "/api/leads/service-requested-options",
        headers=owner_headers,
        params={"page": 2, "size": 2},
    )
    assert second_options.status_code == 200
    assert second_options.json()["data"]["items"] == [
        {"label": "Website Design", "value": "website design"}
    ]
    searched_options = client.get(
        "/api/leads/service-requested-options",
        headers=owner_headers,
        params={"search": "WEB"},
    )
    assert searched_options.status_code == 200
    assert searched_options.json()["data"]["total"] == 1

    staff_headers = login_headers(client, staff)
    staff_options = client.get(
        "/api/leads/service-requested-options",
        headers=staff_headers,
    )
    assert staff_options.status_code == 200
    assert staff_options.json()["data"] == {
        "items": [
            {"label": "SEO", "value": "seo"},
            {"label": "Website Design", "value": "website design"},
        ],
        "page": 1,
        "size": 50,
        "total": 2,
    }
    staff_filtered = client.get(
        "/api/leads",
        headers=staff_headers,
        params={"service_requested": "website design"},
    )
    assert staff_filtered.status_code == 200
    assert staff_filtered.json()["data"]["total"] == 1
    assert staff_filtered.json()["data"]["items"][0]["id"] == website["id"]

    other_tenant = create_tenant(database_session, "service-other-tenant")
    other_owner = create_user(database_session, "service.other.owner")
    add_membership(database_session, other_owner, other_tenant, roles[RoleKey.OWNER])
    other_headers = login_headers(client, other_owner)
    other_create = client.post(
        "/api/leads",
        headers=other_headers,
        json=lead_payload(
            "secret-service@example.com",
            service_requested="Secret Service",
        ),
    )
    assert other_create.status_code == 201
    other_options = client.get(
        "/api/leads/service-requested-options",
        headers=other_headers,
    )
    assert other_options.status_code == 200
    assert other_options.json()["data"]["items"] == [
        {"label": "Secret Service", "value": "secret service"}
    ]
    assert all(
        item["value"] != "secret service"
        for item in first_options.json()["data"]["items"]
    )


def test_downloaded_template_and_service_requested_header_alias_end_to_end(
    client: TestClient,
    database_session: Session,
) -> None:
    _, _, _, headers = owner_setup(client, database_session)

    template_response = client.get("/api/leads/import-template", headers=headers)
    assert template_response.status_code == 200
    assert template_response.headers["cache-control"] == "no-store"
    template = load_workbook(BytesIO(template_response.content), read_only=False)
    template_headers = [cell.value for cell in template["Leads"][1]]
    service_column = template_headers.index("service_requested") + 1
    assert template["Leads"].cell(1, service_column).font.bold is True
    assert template["Leads"].cell(1, service_column).fill.fill_type == "solid"

    sample_headers = [cell.value for cell in template["Sample"][1]]
    sample_values = [cell.value for cell in template["Sample"][2]]
    assert dict(zip(sample_headers, sample_values, strict=True))["service_requested"] == (
        "Website Design"
    )

    alias_workbook = Workbook()
    alias_sheet = alias_workbook.active
    alias_sheet.title = "Leads"
    alias_sheet.append(
        ["first_name", "last_name", "email", "Service Requested"]
    )
    alias_sheet.append(
        ["Alias", "Import", "service-alias@example.com", "Website Design"]
    )
    alias_output = BytesIO()
    alias_workbook.save(alias_output)

    preview = client.post(
        "/api/leads/imports/preview",
        headers=headers,
        files={
            "file": (
                "service-alias.xlsx",
                alias_output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 201
    assert preview.json()["data"]["summary"]["valid"] == 1
    batch_id = uuid.UUID(preview.json()["data"]["id"])
    import_row = database_session.scalar(
        select(LeadImportRow).where(LeadImportRow.batch_id == batch_id)
    )
    assert import_row.normalized_data["service_requested"] == "Website Design"

    commit = client.post(f"/api/leads/imports/{batch_id}/commit", headers=headers)
    assert commit.status_code == 200
    assert commit.json()["data"]["summary"]["inserted"] == 1
    database_session.expire_all()
    imported = database_session.scalar(
        select(Lead).where(Lead.email == "service-alias@example.com")
    )
    assert imported is not None
    assert imported.service_requested == "Website Design"

    detail = client.get(f"/api/leads/{imported.id}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["data"]["service_requested"] == "Website Design"

    options = client.get("/api/leads/service-requested-options", headers=headers)
    assert options.status_code == 200
    assert options.json()["data"]["items"] == [
        {"label": "Website Design", "value": "website design"}
    ]
