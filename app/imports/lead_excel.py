import hashlib
import json
import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from app.core.lead_catalogs import (
    CURRENCY_CODES,
    CompanySize,
    ConsentStatus,
    EmailStatus,
    RevenuePeriod,
)
from app.schemas.lead import LeadFields, validation_errors
from app.utils.lead_normalization import country_catalog_names

CANONICAL_COLUMNS = (
    "first_name",
    "last_name",
    "email",
    "designation",
    "company",
    "country",
    "email_status",
    "linkedin_url",
    "facebook_url",
    "instagram_url",
    "x_url",
    "github_url",
    "phone",
    "city",
    "state",
    "company_website",
    "industry",
    "service_requested",
    "company_size",
    "revenue_amount",
    "revenue_currency",
    "revenue_period",
    "notes",
    "consent_status",
    "consent_evidence",
    "consent_reference",
)
REQUIRED_COLUMNS = frozenset({"first_name", "last_name", "email"})
TEXT_ONLY_COLUMNS = frozenset({"phone"})


class LeadImportFileError(Exception):
    pass


@dataclass(frozen=True)
class ParsedLeadRow:
    row_number: int
    data: dict[str, Any]
    errors: list[dict[str, str]]
    row_hash: str


@dataclass(frozen=True)
class ParsedWorkbook:
    rows: list[ParsedLeadRow]
    mapping: dict[str, str]
    file_sha256: str


def normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().casefold()).strip("_")


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _add_list_validation(sheet: Any, column_index: int, formula: str) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Select a documented value from the list"
    validation.errorTitle = "Invalid value"
    sheet.add_data_validation(validation)
    validation.add(f"{sheet.cell(1, column_index).column_letter}2:{sheet.cell(1, column_index).column_letter}5001")


def build_import_template() -> bytes:
    workbook = Workbook()
    leads = workbook.active
    leads.title = "Leads"
    instructions = workbook.create_sheet("Instructions", 0)
    sample = workbook.create_sheet("Sample")
    catalogs = workbook.create_sheet("Catalogs")

    instructions.append(["PowerLead Lead Import Template"])
    instructions.append(["Enter data only in the Leads sheet. The Sample sheet is guidance and is not imported."])
    instructions.append(["Required columns", "first_name, last_name, email"])
    instructions.append(["Country", "Country name, for example Bangladesh or United States; ISO codes remain accepted"])
    instructions.append(["URLs", "Bare domains such as google.com are accepted and stored as https://google.com"])
    instructions.append(["Revenue", "amount, ISO 4217 currency, and period must be supplied together"])
    instructions.append(["Phone", "Keep the cell formatted as Text to preserve leading zero or +"])
    instructions.append(["Consent", "Granted consent requires evidence or a reference; importing never sends messages"])
    instructions["A1"].font = Font(bold=True, size=14)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in (leads, sample):
        sheet.append(list(CANONICAL_COLUMNS))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(CANONICAL_COLUMNS)).column_letter}1"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for index, field in enumerate(CANONICAL_COLUMNS, start=1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = max(14, len(field) + 2)
        phone_column = CANONICAL_COLUMNS.index("phone") + 1
        for row in range(2, 5002):
            sheet.cell(row, phone_column).number_format = "@"

    sample_values = {
        "first_name": "Sample",
        "last_name": "Person",
        "email": "sample@example.com",
        "company": "Example Company",
        "country": "Bangladesh",
        "email_status": "unknown",
        "phone": "+8801000000000",
        "company_website": "google.com",
        "service_requested": "Website Design",
        "company_size": "11_50",
        "revenue_amount": "100000.00",
        "revenue_currency": "USD",
        "revenue_period": "annual",
        "consent_status": "unknown",
    }
    for column, field in enumerate(CANONICAL_COLUMNS, start=1):
        sample.cell(2, column, _safe_cell(sample_values.get(field, "")))

    catalog_values = {
        "email_status": [item.value for item in EmailStatus],
        "company_size": [item.value for item in CompanySize],
        "revenue_period": [item.value for item in RevenuePeriod],
        "consent_status": [item.value for item in ConsentStatus],
        "country": country_catalog_names(),
        "revenue_currency": sorted(CURRENCY_CODES),
    }
    for column, (name, values) in enumerate(catalog_values.items(), start=1):
        catalogs.cell(1, column, name)
        for row, value in enumerate(values, start=2):
            catalogs.cell(row, column, value)
        target_column = CANONICAL_COLUMNS.index(name) + 1
        letter = catalogs.cell(1, column).column_letter
        _add_list_validation(leads, target_column, f"Catalogs!${letter}$2:${letter}${len(values) + 1}")
    catalogs.sheet_state = "hidden"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_mapping(mapping_json: str | None) -> dict[str, str]:
    if not mapping_json:
        return {}
    try:
        raw = json.loads(mapping_json)
    except json.JSONDecodeError as exception:
        raise LeadImportFileError("Mapping must be a valid JSON object") from exception
    if not isinstance(raw, dict):
        raise LeadImportFileError("Mapping must be a JSON object")
    mapping: dict[str, str] = {}
    for source, target in raw.items():
        normalized_source = normalize_header(source)
        normalized_target = normalize_header(target)
        if normalized_target not in CANONICAL_COLUMNS:
            raise LeadImportFileError(f"Unknown canonical field: {target}")
        if normalized_target in mapping.values():
            raise LeadImportFileError(f"Canonical field mapped more than once: {target}")
        mapping[normalized_source] = normalized_target
    return mapping


def parse_workbook(
    file_bytes: bytes,
    *,
    mapping_json: str | None,
    max_rows: int,
) -> ParsedWorkbook:
    try:
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exception:
        raise LeadImportFileError("The uploaded file is not a valid .xlsx workbook") from exception

    sheet = workbook["Leads"] if "Leads" in workbook.sheetnames else workbook.active
    if sheet.max_row > max_rows + 1:
        raise LeadImportFileError(f"Workbook exceeds the {max_rows} row limit")
    if sheet.max_column > 200:
        raise LeadImportFileError("Workbook exceeds the 200 column limit")
    row_iterator = sheet.iter_rows()
    try:
        header_cells = next(row_iterator)
    except StopIteration as exception:
        raise LeadImportFileError("Workbook has no header row") from exception
    if any(cell.data_type == "f" for cell in header_cells):
        raise LeadImportFileError("Formula cells are not allowed in headers")

    source_headers = [normalize_header(cell.value) if cell.value is not None else "" for cell in header_cells]
    if not any(source_headers):
        raise LeadImportFileError("Workbook has no headers")
    explicit_mapping = parse_mapping(mapping_json)
    mapping = {
        source: explicit_mapping.get(source, source)
        for source in source_headers
        if source and explicit_mapping.get(source, source) in CANONICAL_COLUMNS
    }
    if len(set(mapping.values())) != len(mapping.values()):
        raise LeadImportFileError("Multiple source headers map to the same canonical field")
    missing = REQUIRED_COLUMNS - set(mapping.values())
    if missing:
        raise LeadImportFileError(
            "Missing required mapped columns: " + ", ".join(sorted(missing))
        )

    parsed_rows: list[ParsedLeadRow] = []
    for row_number, cells in enumerate(row_iterator, start=2):
        if len(parsed_rows) >= max_rows:
            raise LeadImportFileError(f"Workbook exceeds the {max_rows} row limit")
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        values: dict[str, Any] = {}
        errors: list[dict[str, str]] = []
        for index, cell in enumerate(cells):
            if index >= len(source_headers):
                break
            canonical = mapping.get(source_headers[index])
            if canonical is None:
                continue
            if cell.data_type == "f":
                errors.append({"field": canonical, "message": "Formula cells are not allowed"})
                continue
            value = cell.value
            if value in (None, ""):
                continue
            if canonical in TEXT_ONLY_COLUMNS and value not in (None, "") and not isinstance(value, str):
                errors.append(
                    {
                        "field": canonical,
                        "message": "Phone must be stored as text to preserve formatting",
                    }
                )
                continue
            values[canonical] = value

        original_email = str(values.get("email", "")).strip()
        if not errors:
            try:
                validated = LeadFields.model_validate(values)
                values = validated.model_dump(mode="json")
                values["email_original"] = original_email
            except ValidationError as exception:
                errors.extend(validation_errors(exception))
        stable_json = json.dumps(values, sort_keys=True, default=str, separators=(",", ":"))
        parsed_rows.append(
            ParsedLeadRow(
                row_number=row_number,
                data=values,
                errors=errors,
                row_hash=hashlib.sha256(stable_json.encode("utf-8")).hexdigest(),
            )
        )
    if not parsed_rows:
        raise LeadImportFileError("Workbook contains no data rows")
    return ParsedWorkbook(
        rows=parsed_rows,
        mapping=mapping,
        file_sha256=hashlib.sha256(file_bytes).hexdigest(),
    )
